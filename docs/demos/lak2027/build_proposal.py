from pathlib import Path
import sys
import hashlib
from copy import deepcopy
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import RGBColor

root = Path(__file__).resolve().parent
source = Path(sys.argv[1])
if hashlib.sha256(source.read_bytes()).hexdigest() != '8db958ada6e96efd2c67d915354b0b8ba3865944bee7b8520cec7d9b16953a59':
    raise ValueError('Official template hash differs; review its changed requirements first')
doc = Document(source)
abstract = (root / 'abstract.txt').read_text().strip()
assert len(abstract.split()) == 200
for child in list(doc._element.body):
    if child.tag != qn('w:sectPr'):
        doc._element.body.remove(child)
# Retain the venue's section layout, headers, footers and JLA style definitions.
# Title uses Word's semantic Title style with the supplied JLA_Title formatting.
title_style = doc.styles['Title'] if 'Title' in doc.styles else doc.styles.add_style('Title', WD_STYLE_TYPE.PARAGRAPH)
base = doc.styles['JLA_Title']
title_style.base_style = base
for tag in ('w:pPr', 'w:rPr'):
    old = title_style.element.find(qn(tag))
    if old is not None:
        title_style.element.remove(old)
    value = base.element.find(qn(tag))
    if value is not None:
        title_style.element.append(deepcopy(value))
title_style.font.color.rgb = RGBColor(0, 0, 0)
doc.add_paragraph('Transparent observation handling in a study analytics demonstration', style='Title')
doc.add_paragraph('Ryan Gomez', style='JLA_Authors')
doc.add_paragraph('Format: Interactive demonstration', style='JLA_Author_Details')
doc.add_paragraph('ABSTRACT: ' + ' '.join(abstract.split()), style='JLA_Abstract')
doc.add_paragraph('Keywords: learning analytics; observation quality; missing data; consent; study review', style='JLA_Keywords')
doc.add_paragraph('Demonstration movie', style='JLA_Heading1')
doc.add_paragraph('Captioned walkthrough of the executed example, 110 seconds, no audio.', style='JLA_Normal')
doc.add_paragraph('https://github.com/vertex-studyAI/vertexED.ai/blob/conference/20261007-pilot-observation-integrity/docs/demos/lak2027/walkthrough.mp4', style='JLA_Normal')
doc.add_paragraph('Demonstration scope', style='JLA_Heading1')
doc.add_paragraph('The local tool exposes source validation, incomplete-session interpretation, withdrawal and aggregate export. Every record is constructed. No learner study or intervention outcome is reported. At the demonstration, attendees can question the decisions and inspect the exact source and output.', style='JLA_Normal')
doc.add_paragraph('Use of artificial intelligence', style='JLA_Heading1')
doc.add_paragraph('ChatGPT assisted with drafting, implementation and verification of this demonstration package. Generated code and text are accompanied by executed software evidence. The presenting author remains responsible for reviewing the contribution, attribution, limitations and final submission.', style='JLA_Normal')
doc.core_properties.title = 'Transparent observation handling in a study analytics demonstration'
doc.core_properties.author = 'Ryan Gomez'
doc.core_properties.subject = 'LAK27 interactive demonstration preparation'
doc.core_properties.comments = ''
doc.save(root / 'LAK27_Demo_Proposal.docx')
print(root / 'LAK27_Demo_Proposal.docx')
